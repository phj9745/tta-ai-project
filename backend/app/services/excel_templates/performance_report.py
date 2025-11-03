from __future__ import annotations

import io
import re
import zipfile
from copy import deepcopy
from dataclasses import dataclass
from typing import Dict, Iterable, List, Mapping, Optional, Sequence, Set

from fastapi import HTTPException

from ..performance_report.models import (
    LinuxSample,
    PerformanceDataset,
    PerformanceOSType,
    WindowsSample,
)

__all__ = ["PerformanceWorkbookPayload", "build_performance_workbook"]


WINDOWS_BASE_SHEET = "Windows #1"
LINUX_BASE_SHEET = "Linux #1"
WINDOWS_PREFIX = "Windows"
LINUX_PREFIX = "Linux"
MAX_SHEET_TITLE_LENGTH = 31
INVALID_SHEET_CHARS = re.compile(r"[:\\/?*\[\]]")
START_ROW = 4

WINDOWS_RANGE_COLUMNS = {"A", "B", "C", "G", "H"}
LINUX_RANGE_COLUMNS = {"A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L", "M", "N", "O", "P", "Q"}


@dataclass
class SheetStats:
    sheet_name: str
    os_type: PerformanceOSType
    end_row: int


@dataclass
class PerformanceWorkbookPayload:
    template_bytes: bytes
    windows_datasets: Sequence[PerformanceDataset]
    linux_datasets: Sequence[PerformanceDataset]


def _sanitize_device_name(name: str) -> str:
    sanitized = INVALID_SHEET_CHARS.sub(" ", name)
    sanitized = sanitized.replace("'", "’")
    sanitized = re.sub(r"\s+", " ", sanitized).strip()
    if not sanitized:
        return ""
    return sanitized[:MAX_SHEET_TITLE_LENGTH]


def _ensure_unique_sheet_title(base: str, used: Set[str]) -> str:
    normalized = base.casefold()
    if normalized not in used:
        return base

    counter = 2
    while True:
        suffix = f" ({counter})"
        available = MAX_SHEET_TITLE_LENGTH - len(suffix)
        truncated = base[:available].rstrip()
        if not truncated:
            truncated = base[:available]
        candidate = f"{truncated}{suffix}"
        normalized_candidate = candidate.casefold()
        if normalized_candidate not in used:
            return candidate
        counter += 1


def _build_sheet_titles(prefix: str, datasets: Sequence[PerformanceDataset]) -> List[str]:
    used: Set[str] = set()
    titles: List[str] = []

    for index, dataset in enumerate(datasets, start=1):
        raw_name = str(dataset.metadata.get("device_name") or "").strip()
        base_title = _sanitize_device_name(raw_name)
        if not base_title:
            base_title = f"{prefix} #{index}"
        unique_title = _ensure_unique_sheet_title(base_title, used)
        used.add(unique_title.casefold())
        titles.append(unique_title)

    return titles


def build_performance_workbook(payload: PerformanceWorkbookPayload) -> bytes:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - optional dependency guard
        raise RuntimeError("openpyxl package is required to build performance workbooks.") from exc

    template_stream = io.BytesIO(payload.template_bytes)
    try:
        workbook = load_workbook(template_stream)
    except Exception as exc:  # pragma: no cover - safety net
        raise HTTPException(status_code=500, detail="성능시험 템플릿을 읽지 못했습니다.") from exc

    manager = _PerformanceWorkbookManager(workbook, payload)
    manager.apply()

    output_stream = io.BytesIO()
    workbook.save(output_stream)
    workbook_bytes = output_stream.getvalue()

    return _adjust_chart_and_formula_ranges(workbook_bytes, manager.sheet_stats)


class _PerformanceWorkbookManager:
    def __init__(self, workbook, payload: PerformanceWorkbookPayload) -> None:
        self._workbook = workbook
        self._payload = payload
        self.sheet_stats: List[SheetStats] = []

    def apply(self) -> None:
        windows_sheets = self._prepare_os_sheets(
            base_name=WINDOWS_BASE_SHEET,
            prefix=WINDOWS_PREFIX,
            datasets=self._payload.windows_datasets,
        )
        linux_sheets = self._prepare_os_sheets(
            base_name=LINUX_BASE_SHEET,
            prefix=LINUX_PREFIX,
            datasets=self._payload.linux_datasets,
        )

        self._remove_unused_templates()

        for dataset, worksheet in zip(self._payload.windows_datasets, windows_sheets):
            end_row = self._populate_windows_sheet(worksheet, dataset)
            self.sheet_stats.append(SheetStats(sheet_name=worksheet.title, os_type=PerformanceOSType.WINDOWS, end_row=end_row))

        for dataset, worksheet in zip(self._payload.linux_datasets, linux_sheets):
            end_row = self._populate_linux_sheet(worksheet, dataset)
            self.sheet_stats.append(SheetStats(sheet_name=worksheet.title, os_type=PerformanceOSType.LINUX, end_row=end_row))

    def _prepare_os_sheets(
        self,
        *,
        base_name: str,
        prefix: str,
        datasets: Sequence[PerformanceDataset],
    ):
        workbook = self._workbook

        if base_name not in workbook.sheetnames:
            if datasets:
                raise HTTPException(status_code=500, detail=f"성능시험 템플릿에서 '{base_name}' 시트를 찾을 수 없습니다.")
            return []

        template_sheet = workbook[base_name]
        template_index = workbook.sheetnames.index(base_name)
        template_charts = [deepcopy(chart) for chart in getattr(template_sheet, "_charts", [])]
        template_images = [deepcopy(image) for image in getattr(template_sheet, "_images", [])]

        for name in list(workbook.sheetnames):
            if name.startswith(f"{prefix} #") and name != base_name:
                workbook.remove(workbook[name])

        if not datasets:
            workbook.remove(template_sheet)
            return []

        new_sheets: List = []
        for _ in range(len(datasets)):
            copied = workbook.copy_worksheet(template_sheet)
            new_sheets.append(copied)

        workbook.remove(template_sheet)

        for offset, sheet in enumerate(new_sheets):
            try:
                workbook._sheets.remove(sheet)  # type: ignore[attr-defined]
            except ValueError:
                pass
            workbook._sheets.insert(template_index + offset, sheet)  # type: ignore[attr-defined]

        sheet_titles = _build_sheet_titles(prefix, datasets)

        for worksheet, title in zip(new_sheets, sheet_titles):
            worksheet.title = title
            worksheet._charts = []  # type: ignore[attr-defined]
            for chart in template_charts:
                new_chart = deepcopy(chart)
                worksheet.add_chart(new_chart, chart.anchor)
            if template_images:
                worksheet._images = [deepcopy(image) for image in template_images]  # type: ignore[attr-defined]

        return new_sheets

    def _remove_unused_templates(self) -> None:
        workbook = self._workbook
        for name in list(workbook.sheetnames):
            if name.startswith("Android #") or name.startswith("iOS #"):
                workbook.remove(workbook[name])

    def _populate_windows_sheet(self, worksheet, dataset: PerformanceDataset) -> int:
        samples = [sample for sample in dataset.require_windows_samples() if isinstance(sample, WindowsSample)]
        if not samples:
            return START_ROW

        for offset, sample in enumerate(samples):
            row = START_ROW + offset
            worksheet.cell(row=row, column=4, value=sample.timestamp.strftime("%Y-%m-%d %H:%M:%S"))
            worksheet.cell(row=row, column=5, value=_safe_round(sample.cpu_percent))
            if sample.private_bytes is not None:
                worksheet.cell(row=row, column=6, value=int(sample.private_bytes))
            else:
                worksheet.cell(row=row, column=6, value=None)

        end_row = START_ROW + len(samples) - 1
        next_row = end_row + 1
        for column in range(4, 7):
            worksheet.cell(row=next_row, column=column, value=None)

        memory_value = dataset.metadata.get("memory_gb")
        if isinstance(memory_value, (int, float)) and memory_value > 0:
            worksheet["M1"] = float(memory_value)
            worksheet["K8"] = float(memory_value)
        else:
            worksheet["M1"] = None
            worksheet["K8"] = None

        device_name = str(dataset.metadata.get("device_name") or "").strip()
        if device_name:
            worksheet["M2"] = device_name
            worksheet["K9"] = device_name
        else:
            worksheet["M2"] = None
            worksheet["K9"] = None

        return end_row

    def _populate_linux_sheet(self, worksheet, dataset: PerformanceDataset) -> int:
        samples = [sample for sample in dataset.require_linux_samples() if isinstance(sample, LinuxSample)]
        if not samples:
            return START_ROW

        total_memory_kib = dataset.metadata.get("total_memory_kib")
        memory_value = dataset.metadata.get("memory_gb")
        if isinstance(memory_value, (int, float)) and memory_value > 0:
            worksheet["V1"] = float(memory_value)
            worksheet["T8"] = float(memory_value)
        elif isinstance(total_memory_kib, (int, float)) and total_memory_kib > 0:
            worksheet["V1"] = round(total_memory_kib / 1024 / 1024, 3)
            worksheet["T8"] = round(total_memory_kib / 1024 / 1024, 3)
        else:
            worksheet["V1"] = None
            worksheet["T8"] = None

        device_name = str(dataset.metadata.get("device_name") or "").strip()
        if device_name:
            worksheet["V2"] = device_name
            worksheet["T9"] = device_name
        else:
            worksheet["V2"] = None
            worksheet["T9"] = None

        for offset, sample in enumerate(samples):
            row = START_ROW + offset
            worksheet.cell(row=row, column=4, value=_safe_round(sample.cpu_user_percent))
            worksheet.cell(row=row, column=5, value=_safe_round(sample.cpu_system_percent))
            worksheet.cell(row=row, column=6, value=sample.free_kib)
            worksheet.cell(row=row, column=7, value=sample.buff_kib)
            worksheet.cell(row=row, column=8, value=sample.cache_kib)
            worksheet.cell(row=row, column=9, value=_safe_round(sample.io_bi))
            worksheet.cell(row=row, column=10, value=_safe_round(sample.io_bo))

        end_row = START_ROW + len(samples) - 1
        next_row = end_row + 1
        for column in range(4, 11):
            worksheet.cell(row=next_row, column=column, value=None)
        return end_row


def _safe_round(value: Optional[float], digits: int = 3) -> Optional[float]:
    if value is None:
        return None
    try:
        return round(float(value), digits)
    except (TypeError, ValueError):
        return None


def _adjust_chart_and_formula_ranges(workbook_bytes: bytes, sheet_stats: Sequence[SheetStats]) -> bytes:
    if not sheet_stats:
        return workbook_bytes

    with zipfile.ZipFile(io.BytesIO(workbook_bytes), "r") as source:
        entries = {info.filename: source.read(info.filename) for info in source.infolist()}

    sheet_map = _build_sheet_path_map(entries)
    sheet_chart_map = _build_sheet_chart_map(entries, sheet_map)

    for stats in sheet_stats:
        sheet_path = sheet_map.get(stats.sheet_name)
        if not sheet_path:
            continue
        columns = WINDOWS_RANGE_COLUMNS if stats.os_type == PerformanceOSType.WINDOWS else LINUX_RANGE_COLUMNS
        base_name = WINDOWS_BASE_SHEET if stats.os_type == PerformanceOSType.WINDOWS else LINUX_BASE_SHEET
        entries[sheet_path] = _rewrite_sheet_ranges(
            entries[sheet_path], columns, stats.end_row, stats.sheet_name, base_name
        )

    for stats in sheet_stats:
        chart_paths = sheet_chart_map.get(stats.sheet_name, set())
        if not chart_paths:
            continue
        columns = WINDOWS_RANGE_COLUMNS if stats.os_type == PerformanceOSType.WINDOWS else LINUX_RANGE_COLUMNS
        base_name = WINDOWS_BASE_SHEET if stats.os_type == PerformanceOSType.WINDOWS else LINUX_BASE_SHEET
        for path in chart_paths:
            content = entries.get(path)
            if content is None:
                continue
            entries[path] = _rewrite_chart_ranges(content, stats.sheet_name, base_name, columns, stats.end_row)

    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for path, data in entries.items():
            target.writestr(path, data)
    return output.getvalue()


def _resolve_relationship_target(base_path: str, target: str) -> str:
    import posixpath

    if not target:
        return ""

    # Absolute targets are rooted at the package root, so strip the leading slash
    if target.startswith("/"):
        return target.lstrip("/")

    base_dir = posixpath.dirname(base_path)
    resolved = posixpath.normpath(posixpath.join(base_dir, target))

    while resolved.startswith("../"):
        resolved = resolved[3:]
        resolved = resolved.lstrip("/")

    return resolved


def _build_sheet_path_map(entries: Mapping[str, bytes]) -> Dict[str, str]:
    workbook_xml = entries.get("xl/workbook.xml")
    rels_xml = entries.get("xl/_rels/workbook.xml.rels")
    if workbook_xml is None or rels_xml is None:
        return {}

    import xml.etree.ElementTree as ET

    ns = {"main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    rel_ns = {"rel": "http://schemas.openxmlformats.org/package/2006/relationships"}
    workbook_tree = ET.fromstring(workbook_xml)
    rels_tree = ET.fromstring(rels_xml)

    rel_map = {rel.attrib["Id"]: rel.attrib["Target"] for rel in rels_tree.findall("rel:Relationship", rel_ns)}

    sheet_map: Dict[str, str] = {}
    for sheet in workbook_tree.findall("main:sheets/main:sheet", ns):
        name = sheet.attrib.get("name")
        rel_id = sheet.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
        if not name or not rel_id:
            continue
        target = _resolve_relationship_target("xl/workbook.xml", rel_map.get(rel_id, ""))
        if not target or not target.startswith("xl/worksheets/"):
            continue
        sheet_map[name] = target
    return sheet_map


def _build_sheet_chart_map(entries: Mapping[str, bytes], sheet_map: Mapping[str, str]) -> Mapping[str, Set[str]]:
    import xml.etree.ElementTree as ET
    import posixpath

    ns = {
        "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
        "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
        "drawing": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
        "chart": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    }

    chart_map: Dict[str, Set[str]] = {}

    for sheet_name, sheet_path in sheet_map.items():
        sheet_dir = posixpath.dirname(sheet_path)
        rels_name = posixpath.join(sheet_dir, "_rels", f"{posixpath.basename(sheet_path)}.rels")
        rels_content = entries.get(rels_name)
        if not rels_content:
            continue

        rels_tree = ET.fromstring(rels_content)
        drawing_targets = {}
        for rel in rels_tree.findall("rel:Relationship", ns):
            rel_type = rel.attrib.get("Type", "")
            if rel_type.endswith("/drawing"):
                rel_id = rel.attrib.get("Id")
                if not rel_id:
                    continue
                target = _resolve_relationship_target(sheet_path, rel.attrib.get("Target", ""))
                if not target:
                    continue
                drawing_targets[rel_id] = target

        if not drawing_targets:
            continue

        sheet_charts: Set[str] = set()

        for rel_id, drawing_path in drawing_targets.items():
            drawing_content = entries.get(drawing_path)
            if not drawing_content:
                continue

            drawing_tree = ET.fromstring(drawing_content)
            drawing_rels_path = posixpath.join(
                posixpath.dirname(drawing_path), "_rels", f"{posixpath.basename(drawing_path)}.rels"
            )
            drawing_rels = entries.get(drawing_rels_path)
            if not drawing_rels:
                continue
            drawing_rels_tree = ET.fromstring(drawing_rels)
            chart_targets = {
                rel.attrib.get("Id"): _resolve_relationship_target(
                    drawing_path, rel.attrib.get("Target", "")
                )
                for rel in drawing_rels_tree.findall("rel:Relationship", ns)
                if rel.attrib.get("Type", "").endswith("/chart") and rel.attrib.get("Id")
            }

            for chart_rel in drawing_tree.findall(
                ".//drawing:graphicFrame/a:graphic/a:graphicData/chart:chart", ns
            ):
                chart_id = chart_rel.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
                if not chart_id:
                    continue
                chart_path = chart_targets.get(chart_id)
                if not chart_path:
                    continue
                sheet_charts.add(chart_path)

        if sheet_charts:
            chart_map[sheet_name] = sheet_charts

    return chart_map


def _rewrite_sheet_ranges(
    content: bytes,
    allowed_columns: Iterable[str],
    end_row: int,
    sheet_name: str,
    base_sheet_name: str,
) -> bytes:
    allowed = set(allowed_columns)
    text = content.decode("utf-8")

    if sheet_name != base_sheet_name:
        escaped_base = re.escape(base_sheet_name)
        pattern_sheet = re.compile(rf"(?:'{escaped_base}'|{escaped_base})!")

        def replace_sheet(match: re.Match[str]) -> str:
            if match.group(0).startswith("'"):
                return f"'{sheet_name}'!"
            return f"{sheet_name}!"

        text = pattern_sheet.sub(replace_sheet, text)

    pattern = re.compile(r"\$([A-Z]+)\$4:\$([A-Z]+)\$(\d+)")

    def replace(match: re.Match[str]) -> str:
        col1, col2, _ = match.groups()
        if col1 == col2 and col1 in allowed:
            return f"${col1}$4:${col1}${end_row}"
        return match.group(0)

    updated = pattern.sub(replace, text)
    return updated.encode("utf-8")


def _rewrite_chart_ranges(
    content: bytes,
    sheet_name: str,
    base_sheet_name: str,
    allowed_columns: Iterable[str],
    end_row: int,
) -> bytes:
    allowed = set(allowed_columns)
    text = content.decode("utf-8")

    if sheet_name != base_sheet_name:
        escaped_base = re.escape(base_sheet_name)
        sheet_pattern = re.compile(rf"(?:'{escaped_base}'|{escaped_base})!")

        def replace_sheet(match: re.Match[str]) -> str:
            if match.group(0).startswith("'"):
                return f"'{sheet_name}'!"
            return f"{sheet_name}!"

        text = sheet_pattern.sub(replace_sheet, text)

    escaped = re.escape(sheet_name)
    pattern = re.compile(rf"(?:'{escaped}'|{escaped})!\$([A-Z]+)\$4:\$([A-Z]+)\$(\d+)")

    def replace(match: re.Match[str]) -> str:
        col1, col2, _ = match.groups()
        if col1 == col2 and col1 in allowed:
            prefix = match.group(0).split("!")[0]
            return f"{prefix}!${col1}$4:${col1}${end_row}"
        return match.group(0)

    updated = pattern.sub(replace, text)
    return updated.encode("utf-8")
