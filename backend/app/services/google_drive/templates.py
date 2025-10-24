"""Template and spreadsheet helpers for Google Drive operations."""
from __future__ import annotations

import io
import logging
import os
import re
import zipfile
from pathlib import Path
from typing import Dict, Mapping, Tuple, TYPE_CHECKING

from fastapi import HTTPException

try:  # pragma: no cover - optional dependency
    from openpyxl import Workbook
except ImportError:  # pragma: no cover
    Workbook = None  # type: ignore[assignment]

from ..excel_templates import defect_report, feature_list, security_report, testcases
from ...token_store import StoredTokens

if TYPE_CHECKING:  # pragma: no cover - type hints only
    from .client import GoogleDriveClient

logger = logging.getLogger(__name__)

__all__ = [
    "TEMPLATE_ROOT",
    "PLACEHOLDER_PATTERNS",
    "SHARED_CRITERIA_FILE_CANDIDATES",
    "PREFERRED_SHARED_CRITERIA_FILE_NAME",
    "FEATURE_LIST_START_ROW",
    "FEATURE_LIST_SHEET_CANDIDATES",
    "SPREADSHEET_RULES",
    "normalize_shared_criteria_name",
    "is_shared_criteria_candidate",
    "build_default_shared_criteria_workbook",
    "load_shared_criteria_template_bytes",
    "replace_placeholders",
    "replace_in_office_document",
    "prepare_template_file_content",
    "guess_mime_type",
    "copy_template_to_drive",
]


TEMPLATE_ROOT = Path(__file__).resolve().parents[3] / "template"
PLACEHOLDER_PATTERNS: Tuple[str, ...] = (
    "GS-B-XX-XXXX",
    "GS-B-2X-XXXX",
    "GS-X-X-XXXX",
)
SHARED_CRITERIA_FILE_CANDIDATES: Tuple[str, ...] = (
    "보안성 결함판단기준표 v1.0.xlsx",
    "결함판단기준표 v1.0.xlsx",
    "결함 판단 기준표 v1.0.xlsx",
    "결함 판단기준표 v1.0.xlsx",
    "공유 결함판단기준표 v1.0.xlsx",
    "공유 결함 판단 기준표 v1.0.xlsx",
)
PREFERRED_SHARED_CRITERIA_FILE_NAME = SHARED_CRITERIA_FILE_CANDIDATES[0]
FEATURE_LIST_START_ROW = 8
FEATURE_LIST_SHEET_CANDIDATES: Tuple[str, ...] = (
    "기능리스트",
    "기능 리스트",
    "feature list",
)

SPREADSHEET_RULES: Dict[str, Mapping[str, object]] = {
    "feature-list": {
        "folder_name": "가.계획",
        "file_suffix": "기능리스트 v1.0.xlsx",
        "populate": feature_list.populate_feature_list,
    },
    "testcase-generation": {
        "folder_name": "나.설계",
        "file_suffix": "테스트케이스.xlsx",
        "populate": testcases.populate_testcase_list,
    },
    "defect-report": {
        "folder_name": "다.수행",
        "file_suffix": "결함리포트 v1.0.xlsx",
        "populate": defect_report.populate_defect_report,
    },
    "security-report": {
        "folder_name": "다.수행",
        "file_suffix": "결함리포트 v1.0.xlsx",
        "populate": security_report.populate_security_report,
    },
}


def normalize_shared_criteria_name(name: str) -> str:
    base = name.strip().lower()
    if base.endswith(".xlsx"):
        base = base[:-5]
    return re.sub(r"\s+", "", base)


_SHARED_CRITERIA_NORMALIZED_NAMES = {
    normalize_shared_criteria_name(candidate)
    for candidate in SHARED_CRITERIA_FILE_CANDIDATES
}


def is_shared_criteria_candidate(filename: str) -> bool:
    try:
        normalized = normalize_shared_criteria_name(filename)
    except Exception:
        return False
    return normalized in _SHARED_CRITERIA_NORMALIZED_NAMES


def build_default_shared_criteria_workbook() -> bytes:
    if Workbook is None:  # pragma: no cover - dependency guard
        raise HTTPException(status_code=500, detail="openpyxl 패키지가 필요합니다.")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "결함판단기준"
    headers = [
        "Invicti 결과",
        "결함 요약",
        "결함정도",
        "발생빈도",
        "품질특성",
        "결함 설명",
        "결함 제외 여부",
    ]
    sheet.append(headers)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def load_shared_criteria_template_bytes() -> bytes:
    for candidate in SHARED_CRITERIA_FILE_CANDIDATES:
        template_path = TEMPLATE_ROOT / candidate
        if template_path.exists():
            return template_path.read_bytes()
    return build_default_shared_criteria_workbook()


def replace_placeholders(text: str, exam_number: str) -> str:
    result = text
    for placeholder in PLACEHOLDER_PATTERNS:
        result = result.replace(placeholder, exam_number)
    return result


def replace_in_office_document(data: bytes, exam_number: str) -> bytes:
    original = io.BytesIO(data)
    updated = io.BytesIO()
    with zipfile.ZipFile(original, "r") as source_zip:
        with zipfile.ZipFile(updated, "w") as target_zip:
            for item in source_zip.infolist():
                content = source_zip.read(item.filename)
                try:
                    decoded = content.decode("utf-8")
                except UnicodeDecodeError:
                    target_zip.writestr(item, content)
                    continue
                replaced = replace_placeholders(decoded, exam_number)
                target_zip.writestr(item, replaced.encode("utf-8"))
    return updated.getvalue()


def prepare_template_file_content(path: Path, exam_number: str) -> bytes:
    raw_bytes = path.read_bytes()
    extension = path.suffix.lower()
    if extension in {".docx", ".xlsx", ".pptx"}:
        raw_bytes = replace_in_office_document(raw_bytes, exam_number)
    return raw_bytes


def guess_mime_type(path: Path) -> str:
    import mimetypes

    mime_type, _ = mimetypes.guess_type(path.name)
    return mime_type or "application/octet-stream"


async def copy_template_to_drive(
    client: "GoogleDriveClient",
    tokens: StoredTokens,
    *,
    parent_id: str,
    exam_number: str,
) -> StoredTokens:
    if not TEMPLATE_ROOT.exists():
        raise HTTPException(status_code=500, detail="template 폴더를 찾을 수 없습니다.")

    path_to_folder_id: Dict[Path, str] = {TEMPLATE_ROOT: parent_id}
    active_tokens = tokens
    for root_dir, dirnames, filenames in os.walk(TEMPLATE_ROOT):
        current_path = Path(root_dir)
        drive_parent_id = path_to_folder_id[current_path]

        for dirname in sorted(dirnames):
            local_dir = current_path / dirname
            folder_name = replace_placeholders(dirname, exam_number)
            folder, active_tokens = await client.create_child_folder(
                active_tokens,
                name=folder_name,
                parent_id=drive_parent_id,
            )
            path_to_folder_id[local_dir] = str(folder["id"])

        for filename in sorted(filenames):
            if is_shared_criteria_candidate(filename):
                logger.info("Skip copying shared criteria into project: %s", filename)
                continue

            local_file = current_path / filename
            target_name = replace_placeholders(filename, exam_number)
            content = prepare_template_file_content(local_file, exam_number)
            mime_type = guess_mime_type(local_file)
            _, active_tokens = await client.upload_file_to_folder(
                active_tokens,
                file_name=target_name,
                parent_id=drive_parent_id,
                content=content,
                content_type=mime_type,
            )

    return active_tokens
