from __future__ import annotations

import io
import sys
from pathlib import Path
from typing import Sequence

import pytest

pytest.importorskip("docx")
pytest.importorskip("pypdf")
from docx import Document

BACKEND_ROOT = Path(__file__).resolve().parents[1]
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

from app.services.google_drive.metadata import (  # noqa: E402
    build_project_folder_name,
    extract_project_metadata,
    normalize_label,
    _extract_product_name,
)


def _build_sample_agreement() -> bytes:
    document = Document()
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "시험 신청 번호"
    table.cell(0, 1).text = "GS-B-12-3456"
    table.cell(1, 0).text = "제조자"
    table.cell(1, 1).text = "Acme Corp"
    table.cell(2, 0).text = "제품명및버전"
    table.cell(2, 1).text = "Wonder Widget\n1.0"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _build_sample_pdf_agreement() -> bytes:
    return _build_pdf_agreement_stream(
        [
            "시험 신청 번호 : GS-B-12-3456",
            "제조자 : Acme Corp",
            "제품명 및 버전 : Wonder Widget 1.0",
        ]
    )


def _build_multiline_product_pdf_agreement() -> bytes:
    return _build_pdf_agreement_stream(
        [
            "시험 신청 번호 : GS-B-12-3456",
            "제조자 : Acme Corp",
            "제품명 및 버전 : Wonder Widget",
            "1.0",
        ]
    )


def _build_multilingual_agreement() -> bytes:
    document = Document()
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "시험 신청 번호"
    table.cell(0, 1).text = "GS-B-12-3456"
    table.cell(1, 0).text = "제조자"
    table.cell(1, 1).text = "Acme Corp"
    table.cell(2, 0).text = "제품명및버전"
    table.cell(2, 1).text = "Wonder Widget 1.0\n중장비 호출 v1.0"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _build_multilingual_pdf_agreement() -> bytes:
    return _build_pdf_agreement_stream(
        [
            "시험 신청 번호 : GS-B-12-3456",
            "제조자 : Acme Corp",
            "제품명 및 버전 : Wonder Widget 1.0",
            "중장비 호출 v1.0",
        ]
    )


def _build_korean_name_only_pdf_agreement() -> bytes:
    return _build_pdf_agreement_stream(
        [
            "시험 신청 번호 : GS-B-12-3456",
            "제조자 : Acme Corp",
            "국문명 : 오픈마루 v1.0",
        ]
    )


def _build_table_style_pdf_agreement() -> bytes:
    return _build_pdf_agreement_stream(
        [
            "시험 신청 번호 : GS-B-12-3456",
            "제조자 : Acme Corp",
            "제품명 ( V ) ( 국문 ) 오픈마루",
            "캅 v1.0",
            "(영문) OPENMARU COP v1.0",
        ]
    )


def _build_multiline_label_pdf_agreement() -> bytes:
    return _build_pdf_agreement_stream(
        [
            "시험 신청 번호 : GS-B-12-3456",
            "제조자 : Acme Corp",
            "제품명 ( V )",
            "(국문) 오픈마루",
            "캅 v1.0",
            "(영문) OPENMARU COP v1.0",
        ]
    )


def _build_pdf_agreement_stream(lines: Sequence[str]) -> bytes:
    stream_segments = ["BT /F1 12 Tf 72 720 Td"]
    for index, line in enumerate(lines):
        prefix = "" if index == 0 else "T* "
        escaped_line = (
            line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        )
        stream_segments.append(f"{prefix}({escaped_line}) Tj")
    stream_segments.append("ET")
    stream_content = " ".join(stream_segments).encode("utf-8")

    buffer = io.BytesIO()
    buffer.write(b"%PDF-1.4\n")
    offsets: list[int] = []

    def _write_object(payload: bytes) -> None:
        offsets.append(buffer.tell())
        buffer.write(payload)
        if not payload.endswith(b"\n"):
            buffer.write(b"\n")

    _write_object(b"1 0 obj<< /Type /Catalog /Pages 2 0 R >>endobj\n")
    _write_object(b"2 0 obj<< /Type /Pages /Kids [3 0 R] /Count 1 >>endobj\n")
    _write_object(
        b"3 0 obj<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>endobj\n"
    )
    stream_header = f"4 0 obj<< /Length {len(stream_content)} >>stream\n".encode("ascii")
    _write_object(stream_header + stream_content + b"\nendstream\nendobj\n")
    _write_object(b"5 0 obj<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>endobj\n")

    xref_offset = buffer.tell()
    buffer.write(f"xref\n0 {len(offsets) + 1}\n".encode("ascii"))
    buffer.write(b"0000000000 65535 f \n")
    for offset in offsets:
        buffer.write(f"{offset:010d} 00000 n \n".encode("ascii"))
    buffer.write(f"trailer<< /Size {len(offsets) + 1} /Root 1 0 R >>\n".encode("ascii"))
    buffer.write(b"startxref\n")
    buffer.write(f"{xref_offset}\n".encode("ascii"))
    buffer.write(b"%%EOF\n")
    return buffer.getvalue()


def test_normalize_label_strips_whitespace() -> None:
    assert normalize_label(" 제 조 자 ") == "제조자"


def test_extract_project_metadata_reads_table() -> None:
    metadata = extract_project_metadata(_build_sample_agreement())
    assert metadata == {
        "exam_number": "GS-B-12-3456",
        "company_name": "Acme Corp",
        "product_name": "Wonder Widget 1.0",
    }


def test_extract_project_metadata_reads_pdf() -> None:
    metadata = extract_project_metadata(_build_sample_pdf_agreement(), file_extension=".pdf")
    assert metadata == {
        "exam_number": "GS-B-12-3456",
        "company_name": "Acme Corp",
        "product_name": "Wonder Widget 1.0",
    }


def test_extract_project_metadata_reads_pdf_with_multiline_product() -> None:
    metadata = extract_project_metadata(
        _build_multiline_product_pdf_agreement(), file_extension=".pdf"
    )
    assert metadata == {
        "exam_number": "GS-B-12-3456",
        "company_name": "Acme Corp",
        "product_name": "Wonder Widget 1.0",
    }


def test_extract_project_metadata_prefers_korean_product_name_from_docx() -> None:
    metadata = extract_project_metadata(_build_multilingual_agreement())
    assert metadata == {
        "exam_number": "GS-B-12-3456",
        "company_name": "Acme Corp",
        "product_name": "중장비 호출 v1.0",
    }


def test_extract_project_metadata_prefers_korean_product_name_from_pdf() -> None:
    metadata = extract_project_metadata(
        _build_multilingual_pdf_agreement(), file_extension=".pdf"
    )
    assert metadata == {
        "exam_number": "GS-B-12-3456",
        "company_name": "Acme Corp",
        "product_name": "중장비 호출 v1.0",
    }


def test_extract_project_metadata_reads_korean_name_label_from_pdf() -> None:
    metadata = extract_project_metadata(
        _build_korean_name_only_pdf_agreement(), file_extension=".pdf"
    )
    assert metadata == {
        "exam_number": "GS-B-12-3456",
        "company_name": "Acme Corp",
        "product_name": "오픈마루 v1.0",
    }


def test_extract_product_name_strips_english_suffix_from_mixed_line() -> None:
    assert (
        _extract_product_name("오픈마루 캅 v1.0(영문) OPENMARU COP v1.0")
        == "오픈마루 캅 v1.0"
    )


def test_extract_project_metadata_reads_table_style_product_rows_from_pdf() -> None:
    metadata = extract_project_metadata(
        _build_table_style_pdf_agreement(), file_extension=".pdf"
    )
    assert metadata == {
        "exam_number": "GS-B-12-3456",
        "company_name": "Acme Corp",
        # The minimal PDF fixture truncates the trailing "캅" glyph, but we still
        # verify that the Korean line is preferred over the English continuation.
        "product_name": "오픈마루 v1.0",
    }


def test_extract_project_metadata_reads_multiline_product_label_from_pdf() -> None:
    metadata = extract_project_metadata(
        _build_multiline_label_pdf_agreement(), file_extension=".pdf"
    )
    assert metadata == {
        "exam_number": "GS-B-12-3456",
        "company_name": "Acme Corp",
        "product_name": "오픈마루 v1.0",
    }


def test_build_project_folder_name_formats_metadata() -> None:
    metadata = {
        "exam_number": "GS-B-12-3456",
        "company_name": "Acme Corp",
        "product_name": "Wonder Widget 1.0",
    }
    assert build_project_folder_name(metadata) == "[GS-B-12-3456] Acme Corp - Wonder Widget 1.0"
